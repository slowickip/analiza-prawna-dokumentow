from __future__ import annotations

import hashlib
import io
import logging
import shutil
import subprocess
import zipfile
from collections.abc import Callable
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from docx import Document
from fpdf import FPDF
from PIL import Image, ImageDraw, ImageFont

from contract_analyzer.config import RunConfig
from contract_analyzer.domain import ReadMode, ReferenceRecord, SourceAnchor, Unit
from contract_analyzer.ingest import IngestError, IngestService
from contract_analyzer.ingest.doc_conversion import _validate_doc_conversion_output
from contract_analyzer.ingest.ocr import _page_ink_share, _read_pdf_ocr
from contract_analyzer.ingest.pdf import _select_pdf_mode
from font_paths import resolve_test_font

FIXTURES = Path(__file__).resolve().parent.parent / "fixtures"
OLE2_HEADER = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
# Hostile ZIPs are built while parametrization is evaluated, so their bytes end up in
# pytest node IDs. ZipFile.writestr() stamps entries with the local clock, so the bytes
# differ between xdist workers that cross a DOS timestamp tick and collection then fails
# on mismatched IDs. Every entry gets this fixed DOS-epoch instant instead.
ZIP_FIXED_DATE_TIME = (1980, 1, 1, 0, 0, 0)


@pytest.fixture
def service() -> IngestService:
    return IngestService()


@pytest.fixture
def fixture_bytes() -> object:
    def _load(name: str) -> bytes:
        return (FIXTURES / name).read_bytes()

    return _load


def _mixed_two_page_pdf() -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="UMOWA NAJMU")
    pdf.add_page()
    return bytes(pdf.output())


def _native_plus_image_page_pdf() -> bytes:
    from pypdf import PdfReader, PdfWriter

    native = FPDF()
    native.add_page()
    native.set_font("Helvetica", size=12)
    native.cell(text="UMOWA NAJMU")
    font = ImageFont.truetype(resolve_test_font(), 18)
    scanned_page = Image.new("RGB", (800, 400), "white")
    ImageDraw.Draw(scanned_page).text(
        (40, 40), "SKANOWANA STRONA", fill="black", font=font
    )
    scanned_buffer = io.BytesIO()
    scanned_page.save(scanned_buffer, format="PDF", resolution=150.0)
    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(bytes(native.output()))))
    writer.append(PdfReader(scanned_buffer))
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def _three_page_scanned_with_blank_middle() -> bytes:
    from pypdf import PdfReader, PdfWriter

    font = ImageFont.truetype(resolve_test_font(), 18)

    def page_bytes(text: str | None) -> bytes:
        image = Image.new("RGB", (800, 400), "white")
        if text is not None:
            ImageDraw.Draw(image).text((40, 40), text, fill="black", font=font)
        buffer = io.BytesIO()
        image.save(buffer, format="PDF", resolution=150.0)
        return buffer.getvalue()

    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(page_bytes("UMOWA STRONA 1"))))
    writer.append(PdfReader(io.BytesIO(page_bytes(None))))
    writer.append(PdfReader(io.BytesIO(page_bytes("UMOWA STRONA 3"))))
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def _native_three_page_with_blank_middle() -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="STRONA NATYWNA JEDEN")
    pdf.add_page()
    pdf.add_page()
    pdf.cell(text="STRONA NATYWNA TRZY")
    return bytes(pdf.output())


def _vertical_merge_docx() -> bytes:
    document = Document()
    table = document.add_table(rows=2, cols=1)
    merged = table.rows[0].cells[0]
    merged.merge(table.rows[1].cells[0])
    merged.text = "VERTICAL CLAUSE"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _nested_order_docx() -> bytes:
    document = Document()
    table = document.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cell.paragraphs[0].text = "BEFORE"
    nested = cell.add_table(rows=1, cols=1)
    nested.rows[0].cells[0].text = "NESTED"
    cell.add_paragraph("AFTER")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _fixed_timestamp_entry(name: str) -> zipfile.ZipInfo:
    """What writestr() builds for a plain name, minus its dependency on the clock."""
    info = zipfile.ZipInfo(name, date_time=ZIP_FIXED_DATE_TIME)
    info.compress_type = zipfile.ZIP_STORED
    info.external_attr = 0o600 << 16
    return info


def _malformed_docx_zip() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            _fixed_timestamp_entry("[Content_Types].xml"),
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" '
            'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument'
            '.wordprocessingml.document.main+xml"/>'
            "</Types>",
        )
        archive.writestr(
            _fixed_timestamp_entry("_rels/.rels"),
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            'relationships/officeDocument" '
            'Target="word/document.xml"/>'
            "</Relationships>",
        )
        archive.writestr(
            _fixed_timestamp_entry("word/document.xml"), "<w:document><unclosed"
        )
        archive.writestr(
            _fixed_timestamp_entry("word/_rels/document.xml.rels"),
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>',
        )
    return buffer.getvalue()


def _empty_zip() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w"):
        pass
    return buffer.getvalue()


def _zip_wrong_parts() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(_fixed_timestamp_entry("wrong.txt"), "not office xml")
    return buffer.getvalue()


def _ooxml_no_rels() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            _fixed_timestamp_entry("[Content_Types].xml"),
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
        )
        archive.writestr(
            _fixed_timestamp_entry("word/document.xml"),
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            "<w:body><w:p><w:r><w:t>x</w:t></w:r></w:p></w:body></w:document>",
        )
    return buffer.getvalue()


def _nested_table_docx() -> bytes:
    document = Document()
    document.add_paragraph("VISIBLE")
    table = document.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cell.text = ""
    nested = cell.add_table(rows=1, cols=1)
    nested.rows[0].cells[0].text = "HIDDEN CLAUSE"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _merged_cell_docx() -> bytes:
    document = Document()
    table = document.add_table(rows=1, cols=2)
    merged = table.rows[0].cells[0]
    merged.merge(table.rows[0].cells[1])
    merged.text = "ONE CLAUSE"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- format readers (TXT, DOCX, DOC, PDF native, PDF OCR) ---


@pytest.mark.parametrize(
    "name", ["sample.txt", "sample.docx", "native.pdf", "scanned.pdf"]
)
def test_supported_fixture_produces_nonempty_canonical_text(
    service: IngestService, fixture_bytes: object, name: str
) -> None:
    document = service.ingest(name, fixture_bytes(name))
    assert document.text.strip()
    assert document.anchors


def test_native_pdf_uses_native_mode(
    service: IngestService, fixture_bytes: object
) -> None:
    document = service.ingest("native.pdf", fixture_bytes("native.pdf"))
    assert document.read_mode is ReadMode.NATIVE_PDF
    assert {anchor.read_mode for anchor in document.anchors} == {ReadMode.NATIVE_PDF}


def test_native_pdf_emits_newlines_and_anchors_slice_to_words(
    service: IngestService, fixture_bytes: object
) -> None:
    document = service.ingest("native.pdf", fixture_bytes("native.pdf"))
    assert "\n" in document.text
    for anchor in document.anchors:
        if anchor.start_offset < anchor.end_offset:
            sliced = document.text[anchor.start_offset : anchor.end_offset]
            assert " " not in sliced
            assert "\n" not in sliced
            assert len(sliced) > 0


def test_native_two_page_pdf_emits_newline_at_page_boundary(
    service: IngestService,
) -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="STRONA JEDEN")
    pdf.add_page()
    pdf.cell(text="STRONA DWA")
    pdf_bytes = bytes(pdf.output())

    document = service.ingest("two_page.pdf", pdf_bytes)
    assert "STRONA JEDEN\nSTRONA DWA" == document.text
    for anchor in document.anchors:
        if anchor.start_offset < anchor.end_offset:
            sliced = document.text[anchor.start_offset : anchor.end_offset]
            assert " " not in sliced
            assert "\n" not in sliced


def test_scanned_pdf_emits_newlines_and_anchors_slice_to_words(
    service: IngestService, fixture_bytes: object
) -> None:
    if shutil.which("tesseract") is None:
        pytest.skip("tesseract not installed")
    document = service.ingest("scanned.pdf", fixture_bytes("scanned.pdf"))
    assert "\n" in document.text
    for anchor in document.anchors:
        if anchor.start_offset < anchor.end_offset:
            sliced = document.text[anchor.start_offset : anchor.end_offset]
            assert " " not in sliced
            assert "\n" not in sliced


def test_pdf_uses_one_read_mode_for_every_page(
    service: IngestService, fixture_bytes: object
) -> None:
    if shutil.which("tesseract") is None:
        pytest.skip("tesseract not installed")
    document = service.ingest("scan.pdf", fixture_bytes("scanned.pdf"))
    assert document.read_mode is ReadMode.OCR_PDF
    assert {anchor.read_mode for anchor in document.anchors} == {ReadMode.OCR_PDF}


def test_scanned_pdf_has_no_native_text_layer(fixture_bytes: object) -> None:
    import pdfplumber

    with pdfplumber.open(FIXTURES / "scanned.pdf") as pdf:
        native_chars = sum(len(page.extract_text() or "") for page in pdf.pages)
    assert native_chars < 10


def test_docx_includes_table_cell_text(
    service: IngestService, fixture_bytes: object
) -> None:
    document = service.ingest("sample.docx", fixture_bytes("sample.docx"))
    assert "Kaucja" in document.text
    assert "9 000 zł" in document.text
    assert any(anchor.block_id is not None for anchor in document.anchors)


def test_nested_docx_table_content_is_included(service: IngestService) -> None:
    document = service.ingest("nested.docx", _nested_table_docx())
    assert "VISIBLE" in document.text
    assert "HIDDEN CLAUSE" in document.text


def test_merged_docx_cells_are_not_duplicated(service: IngestService) -> None:
    document = service.ingest("merged.docx", _merged_cell_docx())
    assert document.text.count("ONE CLAUSE") == 1


def test_vertical_merged_docx_cells_are_not_duplicated(service: IngestService) -> None:
    document = service.ingest("vertical.docx", _vertical_merge_docx())
    assert document.text.count("VERTICAL CLAUSE") == 1


def test_nested_docx_content_follows_document_order(service: IngestService) -> None:
    document = service.ingest("order.docx", _nested_order_docx())
    before = document.text.index("BEFORE")
    nested = document.text.index("NESTED")
    after = document.text.index("AFTER")
    assert before < nested < after


def test_mixed_pdf_recovers_image_only_page_text(fixture_bytes: object) -> None:
    if shutil.which("tesseract") is None:
        pytest.skip("tesseract not installed")
    document = IngestService().ingest("mixed.pdf", fixture_bytes("mixed.pdf"))
    assert document.read_mode is ReadMode.OCR_PDF
    assert "SPECYFICZNY ZAPIS ANEKSU" in document.text.upper()


# --- PDF mode selection, coverage, and blank pages ---


def test_default_coverage_share_sends_mixed_pdf_to_ocr() -> None:
    if shutil.which("tesseract") is None:
        pytest.skip("tesseract not installed")
    mixed = _native_plus_image_page_pdf()
    assert _select_pdf_mode(mixed, RunConfig()) is ReadMode.OCR_PDF
    document = IngestService().ingest("mixed.pdf", mixed)
    assert document.read_mode is ReadMode.OCR_PDF
    assert "SKANOWANA STRONA" in document.text.upper()


def test_lower_coverage_share_selects_native_for_mixed_pdf() -> None:
    mixed = _mixed_two_page_pdf()
    service = IngestService(RunConfig(pdf_native_coverage_share=0.5))
    document = service.ingest("mixed.pdf", mixed)
    assert document.read_mode is ReadMode.NATIVE_PDF


def test_scanned_blank_middle_page_ingests_and_records_blank_page() -> None:
    if shutil.which("tesseract") is None:
        pytest.skip("tesseract not installed")
    document = IngestService().ingest(
        "three-page.pdf", _three_page_scanned_with_blank_middle()
    )
    assert "UMOWA STRONA 1" in document.text
    assert "UMOWA STRONA 3" in document.text
    blank_pages = [
        anchor.page for anchor in document.anchors if anchor.block_id == "blank_page_2"
    ]
    assert blank_pages == [2]


def test_native_pdf_with_blank_page_stays_native() -> None:
    pdf_bytes = _native_three_page_with_blank_middle()
    assert _select_pdf_mode(pdf_bytes, RunConfig()) is ReadMode.NATIVE_PDF
    document = IngestService().ingest("native-blank.pdf", pdf_bytes)
    assert document.read_mode is ReadMode.NATIVE_PDF
    blank_pages = [
        anchor.page for anchor in document.anchors if anchor.block_id == "blank_page_2"
    ]
    assert blank_pages == [2]


# --- anchors and coordinates ---


def test_txt_anchors_cover_canonical_text(
    service: IngestService, fixture_bytes: object
) -> None:
    document = service.ingest("sample.txt", fixture_bytes("sample.txt"))
    assert document.text == "".join(
        document.text[anchor.start_offset : anchor.end_offset]
        for anchor in document.anchors
    )
    assert all(anchor.line_start is not None for anchor in document.anchors)


def test_lone_cr_text_gets_correct_line_numbers(service: IngestService) -> None:
    document = service.ingest("lines.txt", b"first\rsecond\rthird")
    assert [anchor.line_start for anchor in document.anchors] == [1, 2, 3]


@pytest.mark.parametrize(
    ("payload", "expected_lines"),
    [
        (b"a\vb\vc", [1, 2, 3]),
        (b"a\fb", [1, 2]),
        (b"a\x1cb", [1, 2]),
        ("a\u0085b".encode(), [1, 2]),
        ("a\u2028b".encode(), [1, 2]),
    ],
)
def test_txt_line_numbers_advance_for_all_splitline_separators(
    service: IngestService, payload: bytes, expected_lines: list[int]
) -> None:
    document = service.ingest("lines.txt", payload)
    assert [anchor.line_start for anchor in document.anchors] == expected_lines


def test_pdf_anchors_have_page_and_bbox(
    service: IngestService, fixture_bytes: object
) -> None:
    document = service.ingest("native.pdf", fixture_bytes("native.pdf"))
    assert document.anchors
    for anchor in document.anchors:
        assert anchor.page is not None
        assert anchor.bbox is not None
        snippet = document.text[anchor.start_offset : anchor.end_offset]
        assert snippet.strip()


def test_ocr_pdf_anchors_have_page_and_bbox(
    service: IngestService, fixture_bytes: object
) -> None:
    if shutil.which("tesseract") is None:
        pytest.skip("tesseract not installed")
    document = service.ingest("scanned.pdf", fixture_bytes("scanned.pdf"))
    assert "najmu" in document.text.lower() or "umowa" in document.text.lower()
    for anchor in document.anchors:
        assert anchor.page is not None
        assert anchor.bbox is not None


def test_ocr_bbox_stable_across_render_scales(fixture_bytes: object) -> None:
    if not __import__("shutil").which("tesseract"):
        pytest.skip("tesseract not installed")
    data = fixture_bytes("scanned.pdf")
    low_scale = IngestService(RunConfig(pdf_ocr_render_scale=1.0)).ingest(
        "scanned.pdf", data
    )
    high_scale = IngestService(RunConfig(pdf_ocr_render_scale=2.0)).ingest(
        "scanned.pdf", data
    )
    assert low_scale.anchors
    assert high_scale.anchors

    def anchor_text(document: object, anchor: object) -> str:
        return document.text[anchor.start_offset : anchor.end_offset]  # type: ignore[attr-defined]

    low_boxes = {
        anchor_text(low_scale, anchor): anchor.bbox for anchor in low_scale.anchors
    }
    high_boxes = {
        anchor_text(high_scale, anchor): anchor.bbox for anchor in high_scale.anchors
    }
    shared_words = set(low_boxes) & set(high_boxes)
    assert shared_words
    sample_word = max(shared_words, key=len)
    low_bbox = low_boxes[sample_word]
    high_bbox = high_boxes[sample_word]
    assert low_bbox is not None
    assert high_bbox is not None
    for low_value, high_value in zip(low_bbox, high_bbox, strict=True):
        assert low_value == pytest.approx(high_value, rel=0.05, abs=2.0)


def test_content_hash_is_stable(service: IngestService, fixture_bytes: object) -> None:
    first = service.ingest("sample.txt", fixture_bytes("sample.txt"))
    second = service.ingest("sample.txt", fixture_bytes("sample.txt"))
    assert first.content_hash == second.content_hash
    assert first.document_id != second.document_id


# --- typed failures ---
# unsupported, empty, corrupt, missing tool, conversion, OCR, limit breach


@pytest.mark.parametrize(
    "factory", [_zip_wrong_parts, _malformed_docx_zip, _ooxml_no_rels]
)
def test_hostile_zip_factories_stamp_one_fixed_timestamp(
    factory: Callable[[], bytes],
) -> None:
    """Their bytes are parametrization data, so they must not depend on the clock."""
    with zipfile.ZipFile(io.BytesIO(factory())) as archive:
        stamps = {info.date_time for info in archive.infolist()}
    assert stamps == {ZIP_FIXED_DATE_TIME}


@pytest.mark.parametrize(
    ("name", "data"),
    [
        ("truncated.txt", b" "),
        ("truncated.docx", b"PK\x03\x04"),
        ("truncated.pdf", b"%PDF-"),
        ("random.txt", b"\t\n  \t"),
        ("random.docx", bytes((index * 41) % 256 for index in range(128))),
        ("random.pdf", bytes((index * 67) % 256 for index in range(128))),
        ("wrong-parts.docx", _zip_wrong_parts()),
        ("empty.docx", _empty_zip()),
        ("pdf-header.pdf", b"%PDF-1.4\n"),
        ("bad-xml.docx", _malformed_docx_zip()),
        ("no-rels.docx", _ooxml_no_rels()),
    ],
)
def test_hostile_input_only_raises_ingest_error(
    service: IngestService, name: str, data: bytes
) -> None:
    with pytest.raises(IngestError):
        service.ingest(name, data)


def test_misnamed_ooxml_reads_as_docx(service: IngestService) -> None:
    """A .docx served or saved as .doc is read, not called corrupt."""
    document = Document()
    document.add_paragraph("UMOWA NAJMU")
    buffer = io.BytesIO()
    document.save(buffer)

    payload = service.ingest("wezwanie.doc", buffer.getvalue())

    assert payload.read_mode is None
    assert payload.conversion is None
    assert "UMOWA NAJMU" in payload.text


def test_cross_format_misnames_stay_rejected(service: IngestService) -> None:
    """Only the OOXML-under-.doc pair may widen; every other pair stays as it was.

    These three hold on main as well; they are here so the narrow exception above cannot
    quietly grow into general content sniffing. openapi.yaml states that the filename
    chooses the format, and the UI picks its preview from the filename, so a document
    accepted under a foreign extension would render against the wrong viewer.
    """
    document = Document()
    document.add_paragraph("TEXT")
    buffer = io.BytesIO()
    document.save(buffer)
    docx_bytes = buffer.getvalue()

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=12)
    pdf.cell(40, 10, "TEXT")
    pdf_bytes = bytes(pdf.output())

    for name, data in (
        ("contract.pdf", docx_bytes),
        ("contract.docx", pdf_bytes),
        ("contract.docx", OLE2_HEADER),
    ):
        with pytest.raises(IngestError) as exc_info:
            service.ingest(name, data)
        assert exc_info.value.code == "corrupt_input", (name, exc_info.value.code)


@pytest.mark.parametrize(
    ("name", "data"),
    [
        ("truncated.doc", b"PK\x03\x04"),
        ("random.doc", bytes((index * 59) % 256 for index in range(128))),
    ],
)
def test_corrupt_doc_container_raises_corrupt_input(
    service: IngestService, name: str, data: bytes
) -> None:
    with pytest.raises(IngestError) as exc_info:
        service.ingest(name, data)
    assert exc_info.value.code == "corrupt_input"


def test_hostile_doc_input_only_raises_ingest_error(
    service: IngestService, tmp_path: Path
) -> None:
    cases = (
        ("empty.doc", _empty_zip()),
        ("wrong-parts.doc", _zip_wrong_parts()),
        ("no-rels.doc", _ooxml_no_rels()),
    )

    for name, payload in cases:

        def make_fake_run(docx_bytes: bytes) -> object:
            def fake_run(
                args: list[str], **kwargs: object
            ) -> subprocess.CompletedProcess[str]:
                if args[1:2] == ["--version"]:
                    return subprocess.CompletedProcess(
                        args, 0, stdout="LibreOffice 7.6.0.0\n", stderr=""
                    )
                output_dir = Path(args[args.index("--outdir") + 1])
                (output_dir / "input.docx").write_bytes(docx_bytes)
                return subprocess.CompletedProcess(args, 0, stdout="", stderr="")

            return fake_run

        with (
            patch(
                "contract_analyzer.ingest.doc_conversion.shutil.which",
                return_value="/usr/bin/soffice",
            ),
            patch(
                "contract_analyzer.ingest.doc_conversion.tempfile.TemporaryDirectory"
            ) as tmp_mock,
            patch(
                "contract_analyzer.ingest.doc_conversion.subprocess.run",
                side_effect=make_fake_run(payload),
            ),
        ):
            tmp_mock.return_value.__enter__.return_value = str(tmp_path)
            tmp_mock.return_value.__exit__.return_value = False
            with pytest.raises(IngestError):
                service.ingest(name, OLE2_HEADER)


@pytest.mark.parametrize(
    ("name", "data", "code"),
    [
        ("x.rtf", b"x", "unsupported_input"),
        ("x.txt", b"", "empty_input"),
        ("x.txt", b"   \n\t", "empty_input"),
        ("broken.docx", b"not a zip archive", "corrupt_input"),
        ("broken.pdf", b"not a pdf", "corrupt_input"),
        ("no-rels.docx", _ooxml_no_rels(), "corrupt_input"),
    ],
)
def test_ingest_preserves_specific_error_codes(
    service: IngestService, name: str, data: bytes, code: str
) -> None:
    with pytest.raises(IngestError) as exc_info:
        service.ingest(name, data)
    assert exc_info.value.code == code


def test_well_formed_doc_header_reaches_missing_tool_check(
    service: IngestService,
) -> None:
    with patch(
        "contract_analyzer.ingest.doc_conversion.shutil.which", return_value=None
    ):
        with pytest.raises(IngestError) as exc_info:
            service.ingest("legacy.doc", OLE2_HEADER)
    assert exc_info.value.code == "missing_tool"


def test_doc_conversion_output_rejects_replacement_dominated_text() -> None:
    with pytest.raises(IngestError) as exc_info:
        _validate_doc_conversion_output("A" * 89 + "\ufffd" * 11)
    assert exc_info.value.code == "conversion_failed"


def test_doc_conversion_output_allows_isolated_replacement_character() -> None:
    _validate_doc_conversion_output("A" * 99 + "\ufffd")


def test_doc_conversion_output_rejects_disallowed_control_character() -> None:
    with pytest.raises(IngestError) as exc_info:
        _validate_doc_conversion_output("Contract text\x00")
    assert exc_info.value.code == "conversion_failed"


def test_ingest_preserves_conversion_failed_code(
    service: IngestService, tmp_path: Path
) -> None:
    with (
        patch(
            "contract_analyzer.ingest.doc_conversion.shutil.which",
            return_value="/usr/bin/soffice",
        ),
        patch(
            "contract_analyzer.ingest.doc_conversion.tempfile.TemporaryDirectory"
        ) as tmp_mock,
        patch("contract_analyzer.ingest.doc_conversion.subprocess.run") as run_mock,
    ):
        tmp_mock.return_value.__enter__.return_value = str(tmp_path)
        tmp_mock.return_value.__exit__.return_value = False
        run_mock.side_effect = [
            subprocess.CompletedProcess(
                ["soffice", "--version"], 0, stdout="LibreOffice 7.6.0.0\n", stderr=""
            ),
            subprocess.CalledProcessError(1, "soffice"),
        ]
        with pytest.raises(IngestError) as exc_info:
            service.ingest("legacy.doc", OLE2_HEADER)
    assert exc_info.value.code == "conversion_failed"


def test_ingest_preserves_limit_breach_code(fixture_bytes: object) -> None:
    service = IngestService(RunConfig(max_input_bytes=1))
    with pytest.raises(IngestError) as exc_info:
        service.ingest("sample.txt", fixture_bytes("sample.txt"))
    assert exc_info.value.code == "limit_breach"


def test_ingest_preserves_ocr_failed_code(fixture_bytes: object) -> None:
    data = fixture_bytes("scanned.pdf")
    mock_pdf = MagicMock()
    mock_pdf.__len__.return_value = 2
    mock_image = MagicMock()
    mock_pdf.__getitem__.return_value.render.return_value.to_pil.return_value = (
        mock_image
    )
    word = {"text": "UMOWA", "left": "1", "top": "1", "width": "10", "height": "10"}
    with (
        patch("contract_analyzer.ingest.ocr.pdfium.PdfDocument", return_value=mock_pdf),
        patch(
            "contract_analyzer.ingest.ocr._run_tesseract",
            side_effect=[[word], []],
        ),
        patch(
            "contract_analyzer.ingest.ocr._page_ink_share",
            return_value=0.01,
        ),
        patch(
            "contract_analyzer.ingest.pdf._select_pdf_mode",
            return_value=ReadMode.OCR_PDF,
        ),
        patch("contract_analyzer.ingest.pdf._pdf_page_count", return_value=2),
    ):
        with pytest.raises(IngestError) as exc_info:
            IngestService().ingest("two-page.pdf", data)
    assert exc_info.value.code == "ocr_failed"


def test_malformed_docx_xml_raises_typed_error(service: IngestService) -> None:
    with pytest.raises(IngestError, match="corrupt_input"):
        service.ingest("broken.docx", _malformed_docx_zip())


def test_doc_conversion_subprocess_contract(
    service: IngestService, tmp_path: Path
) -> None:
    with (
        patch(
            "contract_analyzer.ingest.doc_conversion.shutil.which",
            return_value="/usr/bin/soffice",
        ),
        patch("contract_analyzer.ingest.doc_conversion.subprocess.run") as run_mock,
        patch(
            "contract_analyzer.ingest.doc_conversion.tempfile.TemporaryDirectory"
        ) as tmp_mock,
    ):
        tmp_mock.return_value.__enter__.return_value = str(tmp_path)
        tmp_mock.return_value.__exit__.return_value = False
        run_mock.side_effect = [
            subprocess.CompletedProcess(
                ["soffice", "--version"], 0, stdout="LibreOffice 7.6.0.0\n", stderr=""
            ),
            subprocess.CalledProcessError(1, "soffice"),
        ]
        with pytest.raises(IngestError, match="conversion_failed"):
            service.ingest("legacy.doc", OLE2_HEADER)
        assert run_mock.call_count == 2
        convert_call = run_mock.call_args_list[1]
        assert convert_call.kwargs["check"] is True
        assert convert_call.kwargs["timeout"] is not None
        assert isinstance(convert_call.args[0], list)
        tmp_mock.return_value.__exit__.assert_called_once()


def test_empty_ocr_page_with_ink_aborts_ingestion(fixture_bytes: object) -> None:
    data = fixture_bytes("scanned.pdf")
    mock_pdf = MagicMock()
    mock_pdf.__len__.return_value = 2
    mock_image = MagicMock()
    mock_pdf.__getitem__.return_value.render.return_value.to_pil.return_value = (
        mock_image
    )

    word = {"text": "UMOWA", "left": "1", "top": "1", "width": "10", "height": "10"}

    with (
        patch("contract_analyzer.ingest.ocr.pdfium.PdfDocument", return_value=mock_pdf),
        patch(
            "contract_analyzer.ingest.ocr._run_tesseract",
            side_effect=[[word], []],
        ),
        patch(
            "contract_analyzer.ingest.ocr._page_ink_share",
            return_value=0.01,
        ),
        patch(
            "contract_analyzer.ingest.pdf._select_pdf_mode",
            return_value=ReadMode.OCR_PDF,
        ),
        patch("contract_analyzer.ingest.pdf._pdf_page_count", return_value=2),
    ):
        with pytest.raises(IngestError, match="ocr_failed") as exc_info:
            IngestService().ingest("two-page.pdf", data)
        assert "page 2" in str(exc_info.value)


def test_isolated_scan_noise_reads_as_blank_page(fixture_bytes: object) -> None:
    import random

    data = fixture_bytes("scanned.pdf")
    noise_image = Image.new("RGB", (2000, 2800), "white")
    pixels = noise_image.load()
    random.seed(0)
    for _ in range(500):
        x = random.randrange(2000)
        y = random.randrange(2800)
        pixels[x, y] = (0, 0, 0)

    mock_pdf = MagicMock()
    mock_pdf.__len__.return_value = 1
    mock_pdf.__getitem__.return_value.render.return_value.to_pil.return_value = (
        noise_image
    )

    with (
        patch("contract_analyzer.ingest.ocr.pdfium.PdfDocument", return_value=mock_pdf),
        patch("contract_analyzer.ingest.ocr._run_tesseract", return_value=[]),
        patch(
            "contract_analyzer.ingest.pdf._select_pdf_mode",
            return_value=ReadMode.OCR_PDF,
        ),
        patch("contract_analyzer.ingest.pdf._pdf_page_count", return_value=1),
    ):
        document = IngestService().ingest("noisy.pdf", data)
    assert any(anchor.block_id == "blank_page_1" for anchor in document.anchors)


def test_blank_render_has_no_ink() -> None:
    assert _page_ink_share(Image.new("RGB", (1240, 1754), "white")) == 0.0


def test_real_line_of_text_has_ink() -> None:
    font = ImageFont.truetype(resolve_test_font(), 24)
    image = Image.new("RGB", (1240, 1754), "white")
    ImageDraw.Draw(image).text((60, 60), "UMOWA NAJMU LOKALU", fill="black", font=font)
    assert _page_ink_share(image) > RunConfig().blank_page_ink_share


def test_isolated_noise_has_sub_threshold_ink_share() -> None:
    import random

    image = Image.new("RGB", (2000, 2800), "white")
    pixels = image.load()
    random.seed(0)
    for _ in range(500):
        x = random.randrange(2000)
        y = random.randrange(2800)
        pixels[x, y] = (0, 0, 0)
    assert _page_ink_share(image) <= RunConfig().blank_page_ink_share


def test_oversized_input_raises_limit_breach_without_subprocess() -> None:
    service = IngestService(RunConfig(max_input_bytes=10))
    with patch("contract_analyzer.ingest.doc_conversion.subprocess.run") as run_mock:
        with pytest.raises(IngestError, match="limit_breach") as exc_info:
            service.ingest("legacy.doc", b"x" * 100)
        assert "max_input_bytes" in str(exc_info.value)
        run_mock.assert_not_called()


def test_pdf_page_limit_raises_before_rasterisation() -> None:
    service = IngestService(RunConfig(max_pdf_pages=1))
    with patch("contract_analyzer.ingest.pdf.pdfium.PdfDocument") as pdf_mock:
        with pytest.raises(IngestError, match="limit_breach") as exc_info:
            service.ingest("mixed.pdf", _mixed_two_page_pdf())
        assert "max_pdf_pages" in str(exc_info.value)
        pdf_mock.assert_not_called()


# --- configuration is honoured (no dead parameters) ---


def test_configured_max_input_bytes_is_honoured(fixture_bytes: object) -> None:
    service = IngestService(RunConfig(max_input_bytes=1))
    with pytest.raises(IngestError, match="limit_breach") as exc_info:
        service.ingest("sample.txt", fixture_bytes("sample.txt"))
    assert "max_input_bytes" in str(exc_info.value)


def test_configured_ocr_language_reaches_tesseract(
    fixture_bytes: object,
) -> None:
    if shutil.which("tesseract") is None:
        pytest.skip("tesseract not installed")

    def fake_run(args: list[str], **kwargs: object) -> subprocess.CompletedProcess[str]:
        return subprocess.CompletedProcess(
            args,
            0,
            stdout=(
                "level\tpage\tblock\tpar\tline\tword\tleft\ttop\twidth\theight\tconf\ttext\n"
                "5\t1\t0\t0\t0\t0\t1\t1\t10\t10\t90\tUMOWA\n"
            ),
            stderr="",
        )

    with patch(
        "contract_analyzer.ingest.ocr.subprocess.run", side_effect=fake_run
    ) as run_mock:
        IngestService(RunConfig(ocr_language="eng")).ingest(
            "scanned.pdf", fixture_bytes("scanned.pdf")
        )
    call_args = run_mock.call_args.args[0]
    assert "-l" in call_args
    assert call_args[call_args.index("-l") + 1] == "eng"


# --- privacy (no text in repr, no content in errors, temp cleanup) ---


def test_document_payload_repr_omits_text(service: IngestService) -> None:
    secret = "TAJNA KLAUZULA UMOWY"
    payload = service.ingest("clause.txt", secret.encode("utf-8"))
    rendered = repr(payload)
    assert secret not in rendered
    assert str(payload.document_id) in rendered


def test_unit_repr_omits_text() -> None:
    secret = "TAJNA KLAUZULA JEDNOSTKI"
    unit = Unit(
        id="u1",
        text=secret,
        anchor=SourceAnchor(start_offset=0, end_offset=len(secret)),
    )
    rendered = repr(unit)
    assert secret not in rendered
    assert "u1" in rendered


def test_reference_record_repr_omits_raw_text() -> None:
    from contract_analyzer.domain import ReferenceStatus, ReferenceType

    secret = "TAJNA KLAUZULA ODNIESIENIA"
    record = ReferenceRecord(
        citing_unit_id="u1",
        reference_type=ReferenceType.EXTERNAL_ACT,
        status=ReferenceStatus.EXTERNAL_ACT,
        raw_text=secret,
    )
    rendered = repr(record)
    assert secret not in rendered
    assert "u1" in rendered


def test_conversion_failure_does_not_chain_captured_output(tmp_path: Path) -> None:
    service = IngestService()
    with (
        patch(
            "contract_analyzer.ingest.doc_conversion.shutil.which",
            return_value="/usr/bin/soffice",
        ),
        patch(
            "contract_analyzer.ingest.doc_conversion.tempfile.TemporaryDirectory"
        ) as tmp_mock,
        patch("contract_analyzer.ingest.doc_conversion.subprocess.run") as run_mock,
    ):
        tmp_mock.return_value.__enter__.return_value = str(tmp_path)
        tmp_mock.return_value.__exit__.return_value = False
        run_mock.side_effect = subprocess.CalledProcessError(
            1, "soffice", output="TAJNA KLAUZULA", stderr=""
        )
        with pytest.raises(IngestError, match="conversion_failed") as exc_info:
            service.ingest("legacy.doc", OLE2_HEADER)
        assert exc_info.value.__cause__ is None
        assert "TAJNA KLAUZULA" not in str(exc_info.value)


def test_tesseract_failure_does_not_chain_captured_output(
    fixture_bytes: object,
) -> None:
    if not Path("/usr/bin/tesseract").exists() and not __import__("shutil").which(
        "tesseract"
    ):
        pytest.skip("tesseract not installed")
    service = IngestService()
    with patch("contract_analyzer.ingest.ocr.subprocess.run") as run_mock:
        run_mock.side_effect = subprocess.CalledProcessError(
            1,
            "tesseract",
            output="TAJNA KLAUZULA",
            stderr="",
        )
        with pytest.raises(IngestError, match="ocr_failed") as exc_info:
            service.ingest("scanned.pdf", fixture_bytes("scanned.pdf"))
        assert exc_info.value.__cause__ is None
        assert "TAJNA KLAUZULA" not in str(exc_info.value)


def test_ocr_temp_image_removed_when_save_fails(fixture_bytes: object) -> None:
    data = fixture_bytes("scanned.pdf")
    config = RunConfig()
    mock_pdf = MagicMock()
    mock_pdf.__len__.return_value = 1
    mock_image = MagicMock()
    mock_image.save.side_effect = OSError("write failed")
    mock_pdf.__getitem__.return_value.render.return_value.to_pil.return_value = (
        mock_image
    )

    handle = MagicMock()
    handle.name = "/tmp/fake-ocr-page.png"
    handle.__enter__ = MagicMock(return_value=handle)
    handle.__exit__ = MagicMock(return_value=False)

    with (
        patch("contract_analyzer.ingest.ocr.pdfium.PdfDocument", return_value=mock_pdf),
        patch(
            "contract_analyzer.ingest.ocr.tempfile.NamedTemporaryFile",
            return_value=handle,
        ),
        patch.object(Path, "unlink") as unlink_mock,
    ):
        with pytest.raises(IngestError, match="ocr_failed"):
            _read_pdf_ocr(data, ReadMode.OCR_PDF, config)
        unlink_mock.assert_called_once_with(missing_ok=True)


# --- conversion provenance ---


def test_doc_conversion_records_provenance(tmp_path: Path) -> None:
    docx_bytes = Document()
    docx_bytes.add_paragraph("Skonwertowana klauzula")
    buffer = io.BytesIO()
    docx_bytes.save(buffer)
    converted = buffer.getvalue()
    expected_hash = hashlib.sha256(converted).hexdigest()

    def fake_run(args: list[str], **kwargs: object) -> subprocess.CompletedProcess[str]:
        if args[1:2] == ["--version"]:
            return subprocess.CompletedProcess(
                args, 0, stdout="LibreOffice 7.6.0.0\n", stderr=""
            )
        output_dir = Path(args[args.index("--outdir") + 1])
        (output_dir / "input.docx").write_bytes(converted)
        return subprocess.CompletedProcess(args, 0, stdout="", stderr="")

    with (
        patch(
            "contract_analyzer.ingest.doc_conversion.shutil.which",
            return_value="/usr/bin/soffice",
        ),
        patch(
            "contract_analyzer.ingest.doc_conversion.tempfile.TemporaryDirectory"
        ) as tmp_mock,
        patch(
            "contract_analyzer.ingest.doc_conversion.subprocess.run",
            side_effect=fake_run,
        ),
    ):
        tmp_mock.return_value.__enter__.return_value = str(tmp_path)
        tmp_mock.return_value.__exit__.return_value = False
        document = IngestService().ingest("legacy.doc", OLE2_HEADER)

    assert document.conversion is not None
    assert document.conversion.converter == "libreoffice"
    assert document.conversion.converter_version.startswith("LibreOffice")
    assert document.conversion.output_hash == expected_hash
    assert "Skonwertowana klauzula" in document.text


def test_ingest_logs_warning_on_byte_limit_breach(
    caplog: pytest.LogCaptureFixture,
) -> None:
    service = IngestService(RunConfig(max_input_bytes=10))
    with caplog.at_level(logging.WARNING):
        with pytest.raises(IngestError, match="limit_breach"):
            service.ingest("TAJNA_KLAUZULA_W_NAZWIE.txt", b"x" * 20)
    assert any("max_input_bytes" in record.message for record in caplog.records)
    assert any(record.levelno == logging.WARNING for record in caplog.records)
    assert "TAJNA_KLAUZULA" not in caplog.text


def test_ingest_logs_warning_on_page_limit_breach(
    caplog: pytest.LogCaptureFixture,
) -> None:
    service = IngestService(RunConfig(max_pdf_pages=1))
    with caplog.at_level(logging.WARNING):
        with pytest.raises(IngestError, match="limit_breach"):
            service.ingest("mixed.pdf", _mixed_two_page_pdf())
    assert any("max_pdf_pages" in record.message for record in caplog.records)


def test_ingest_logs_warning_on_unsupported_extension(
    caplog: pytest.LogCaptureFixture,
) -> None:
    service = IngestService()
    with caplog.at_level(logging.WARNING):
        with pytest.raises(IngestError, match="unsupported_input"):
            service.ingest("umowa.TAJNA_KLAUZULA", b"a,b,c\n1,2,3")
    assert any("unsupported extension" in record.message for record in caplog.records)
    assert "TAJNA_KLAUZULA" not in caplog.text.upper()


def test_ingest_logs_warning_on_txt_decode_failure(
    caplog: pytest.LogCaptureFixture,
) -> None:
    service = IngestService()
    with caplog.at_level(logging.WARNING):
        with pytest.raises(IngestError, match="empty_input"):
            service.ingest("empty.txt", b"   ")
    assert any("TXT decode failed" in record.message for record in caplog.records)


def _table_then_paragraph_docx() -> bytes:
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "LEFT"
    table.rows[0].cells[1].text = "RIGHT"
    document.add_paragraph("AFTER")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_block_ids_count_every_paragraph_including_cells() -> None:
    """A paragraph after a two-cell table keeps the identifier the reader sees."""
    from contract_analyzer.ingest.docx import _iter_docx_blocks

    document = Document(io.BytesIO(_table_then_paragraph_docx()))
    blocks = _iter_docx_blocks(document)
    assert [block_id for block_id, _ in blocks] == ["t0_r0_c0", "t0_r0_c1", "p3"]
