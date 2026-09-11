"""Regenerate committed ingestion fixtures.

Run: uv run --python 3.12 tests/fixtures/generate_fixtures.py
"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from fpdf import FPDF
from PIL import Image, ImageDraw, ImageFont

FIXTURES = Path(__file__).resolve().parent

FONT_CANDIDATES = (
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
)


def _resolve_font_path() -> str:
    for candidate in FONT_CANDIDATES:
        if Path(candidate).is_file():
            return candidate
    raise SystemExit(
        "no suitable font found for PDF fixture construction; "
        "install DejaVu or Liberation Sans, or run on macOS"
    )


FONT_PATH = _resolve_font_path()

CONTRACT_LINES = [
    "UMOWA NAJMU LOKALU UŻYTKOWEGO",
    "",
    "zawarta w dniu 15 marca 2026 r. w Warszawie",
    "",
    "pomiędzy:",
    "Spółką „Zielony Zakątek Sp. z o.o.” z siedzibą w Warszawie, reprezentowaną przez "
    "Kierownika Działu Nieruchomości — zwaną dalej „Wynajmującym”,",
    "a",
    "Panem Janem Kowalskim, zamieszkałym w Krakowie — zwanym dalej „Najemcą”,",
    "",
    "§ 1. Przedmiot umowy",
    "Wynajmujący oddaje Najemcy do używania lokal użytkowy oznaczony numerem L-12.",
    "",
    "§ 2. Czynsz",
    "Najemca zobowiązuje się uiszczać miesięczny czynsz w wysokości 4 500 zł.",
]

TABLE_ROWS = [
    ("Pozycja", "Kwota", "Termin"),
    ("Czynsz podstawowy", "4 500 zł", "do 10. dnia miesiąca"),
    ("Opłaty eksploatacyjne", "650 zł", "do 15. dnia miesiąca"),
    ("Kaucja", "9 000 zł", "przy podpisaniu umowy"),
]


def write_sample_txt() -> None:
    (FIXTURES / "sample.txt").write_text("\n".join(CONTRACT_LINES), encoding="utf-8")


def write_sample_docx() -> None:
    doc = Document()
    for line in CONTRACT_LINES:
        doc.add_paragraph(line)
    table = doc.add_table(rows=len(TABLE_ROWS), cols=3)
    for row_idx, row in enumerate(TABLE_ROWS):
        for col_idx, cell_text in enumerate(row):
            table.rows[row_idx].cells[col_idx].text = cell_text
    doc.save(FIXTURES / "sample.docx")


def write_native_pdf() -> None:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.add_font("FixtureFont", "", FONT_PATH)
    pdf.set_font("FixtureFont", size=11)
    for line in CONTRACT_LINES:
        if not line:
            pdf.ln(4)
            continue
        pdf.multi_cell(w=0, h=6, text=line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(FIXTURES / "native.pdf")


def write_scanned_pdf() -> None:
    font = ImageFont.truetype(FONT_PATH, 18)
    text = "\n".join(line for line in CONTRACT_LINES if line)
    img = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(img)
    draw.multiline_text((60, 60), text, fill="black", font=font, spacing=8)
    img.save(FIXTURES / "scanned.pdf", "PDF", resolution=150.0)


def write_mixed_pdf() -> None:
    from pypdf import PdfReader, PdfWriter

    native = FPDF()
    native.add_page()
    native.set_font("Helvetica", size=12)
    native.cell(text="STRONA NATYWNA JEDEN")
    native.add_page()
    native.cell(text="STRONA NATYWNA DWA")

    font = ImageFont.truetype(FONT_PATH, 24)
    scanned_page = Image.new("RGB", (800, 400), "white")
    draw = ImageDraw.Draw(scanned_page)
    draw.text((40, 40), "SPECYFICZNY ZAPIS ANEKSU", fill="black", font=font)
    scanned_buffer = io.BytesIO()
    scanned_page.save(scanned_buffer, format="PDF", resolution=150.0)

    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(bytes(native.output()))))
    writer.append(PdfReader(scanned_buffer))
    with (FIXTURES / "mixed.pdf").open("wb") as handle:
        writer.write(handle)


def main() -> None:
    write_sample_txt()
    write_sample_docx()
    write_native_pdf()
    write_scanned_pdf()
    write_mixed_pdf()
    print(f"Wrote fixtures under {FIXTURES}")


if __name__ == "__main__":
    main()
